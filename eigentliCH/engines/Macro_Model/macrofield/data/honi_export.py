"""Reader for the current HoNI export, the source of record for the indicator set.

The export is `HoNI_Export 11-25.xlsx`, modified 2025-11-13, and it is HONI-2026: 15 indicators in three
groups of five, 16 economies, annual 2004 to 2024. See docs/MODEL_SPEC.md section 4 for which HoNI vintage
is current and why, and `config/honi_indicators.yaml` for the column map this reads against.

**The export carries scores, not just raw values.** Each country sheet holds a block of *scored* indicators
on the 1 to 5 axis (rows 2 onwards) and a block of the *raw* values beneath it, in the same year order. That
matters for how this is wired: the published scores are used directly rather than re-derived from the raw
values through `config/honi_bands.yaml`, because those bands come from the superseded 2022 workbook whose 83
indicators do not correspond to these 15. Re-scoring against the wrong bands would be worse than using the
publisher's own scores.

What this module therefore does is read the published per-indicator scores, and hand them to
`macrofield.model.honi` as `IndicatorScore` objects so that the aggregation, the coverage floors, the
direction handling and the four-stage taxonomy all remain the programme's own tested logic. Only the scoring
step is the publisher's.

**The published composite is not reproduced and is not meant to be.** The export's own composite is a
cross-sectional min-max rescaling within each year, so it is a relative ranking that is not comparable
across years (MODEL_SPEC section 4). It is read here as a cross-check and reported alongside, never used as
the composite.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

#: Where the export lives, in preference order. The first is the source of record recorded in
#: config/honi_indicators.yaml; the second is the copy the MATLAB controller reads, which has been observed
#: to be byte-identical. Both are checked so a machine with only one of them still works.
DEFAULT_SEARCH_PATHS = (
    Path(
        r"C:\Users\nicol\Desktop\SIM_NAS\Knowledge_Center\Marianne Work Folder"
        r"\LaTex Documents\HoNI\HoNI_Export 11-25.xlsx"
    ),
    Path(r"C:\Users\nicol\Desktop\SIM_NAS\SIM_Tech\Master_Controller\HoNI_Export.xlsx"),
)

#: The sheet that carries the published composite per economy per year.
COMPOSITE_SHEET = "HoNI Score"

#: First row of the scored block on a country sheet. Row 1 is the header.
FIRST_SCORE_ROW = 2

#: Number of blank rows between the scored block and the raw block beneath it.
BLOCK_GAP = 2


class HoNIExportError(RuntimeError):
    """Raised when the export cannot be read or does not have the expected shape."""


@dataclass
class EconomyReading:
    """One economy's export data.

    Attributes:
        sheet: The sheet it came from.
        years: The years present, ascending.
        scores: Indicator key to year to published score on the 1 to 5 axis.
        raw: Indicator key to year to raw value, where the raw block was present.
        published_dimensions: Dimension key to year to the export's own sub-index.
        published_composite: Year to the export's own composite, from the composite sheet.
        derived: The export's own derived columns, for cross-checking.
        notes: Anything a reader must know.
    """

    sheet: str
    years: list[int]
    scores: dict[str, dict[int, float]] = field(default_factory=dict)
    raw: dict[str, dict[int, float]] = field(default_factory=dict)
    published_dimensions: dict[str, dict[int, float]] = field(default_factory=dict)
    published_composite: dict[int, float] = field(default_factory=dict)
    derived: dict[str, dict[int, float]] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    @property
    def latest_year(self) -> int:
        if not self.years:
            raise HoNIExportError(f"{self.sheet} carries no years")
        return self.years[-1]

    def indicator_values(self, year: int) -> dict[str, float]:
        """The published score per indicator for one year, omitting any that is absent."""
        return {
            key: series[year]
            for key, series in self.scores.items()
            if year in series
        }


@dataclass
class Export:
    """The whole export.

    Attributes:
        path: Where it was read from.
        economies: Sheet name to reading.
        indicator_keys: The indicator keys in column order.
        notes: Anything applying to the whole export.
    """

    path: Path
    economies: dict[str, EconomyReading] = field(default_factory=dict)
    indicator_keys: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def for_sheet(self, sheet: str) -> EconomyReading:
        if sheet not in self.economies:
            raise HoNIExportError(
                f"the export has no sheet {sheet!r}. Available: {', '.join(sorted(self.economies))}"
            )
        return self.economies[sheet]


def locate_export(search_paths: tuple[Path, ...] = DEFAULT_SEARCH_PATHS) -> Path:
    """Return the first export that exists.

    Raises:
        HoNIExportError: If none of them do, naming every path tried, since a missing NAS mount is the
            normal reason and the reader needs to know which locations were checked.
    """
    for candidate in search_paths:
        if candidate.exists():
            return candidate
    tried = "\n  ".join(str(p) for p in search_paths)
    raise HoNIExportError(
        f"the HoNI export was not found. Paths tried:\n  {tried}\n"
        f"The export is the source of record for the indicator set, so HoNI cannot be scored without it."
    )


def _number(value: Any) -> float | None:
    """Coerce a cell to a float, or None where it is blank or not numeric."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _year(value: Any) -> int | None:
    number = _number(value)
    if number is None:
        return None
    year = int(round(number))
    # The export's year column runs 2004 onwards. A guard against reading a data row as a year.
    return year if 1900 <= year <= 2200 else None


def load_export(
    config,
    path: Path | None = None,
) -> Export:
    """Read the export against the column map in `config/honi_indicators.yaml`.

    Args:
        config: A loaded Config, used for the indicator and dimension column map.
        path: Override the location. Defaults to the first of DEFAULT_SEARCH_PATHS that exists.

    Returns:
        The export.

    Raises:
        HoNIExportError: If the file is missing, unreadable, or its headers do not match the recorded
            column map. A silent mismatch would attribute one indicator's scores to another.
    """
    from openpyxl import load_workbook

    from macrofield.model.honi import load_indicator_definitions

    document = load_indicator_definitions(Path(config.get("honi.indicators_config")))
    definitions = document.get("indicators") or {}
    dimensions = document.get("dimensions") or {}
    if not definitions:
        raise HoNIExportError("config/honi_indicators.yaml carries no indicators")
    if not dimensions:
        raise HoNIExportError("config/honi_indicators.yaml carries no dimensions")

    source = Path(path) if path is not None else locate_export()
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - openpyxl raises several unrelated types
        raise HoNIExportError(f"could not open {source}: {type(error).__name__}: {error}") from error

    export = Export(path=source, indicator_keys=list(definitions))
    try:
        composite = _read_composite_sheet(workbook)
        sheet_names = [
            name
            for name in workbook.sheetnames
            if name not in ("HoNI", COMPOSITE_SHEET)
        ]
        for name in sheet_names:
            reading = _read_country_sheet(workbook[name], name, definitions, dimensions)
            reading.published_composite = composite.get(name, {})
            if not reading.published_composite:
                reading.notes.append(
                    f"the composite sheet carries no column for {name!r}, so the export's own composite "
                    f"is unavailable as a cross-check"
                )
            export.economies[name] = reading
    finally:
        workbook.close()

    export.notes.append(
        "indicator scores are the export's own, on the 1 to 5 axis. They are not re-derived from the raw "
        "values, because config/honi_bands.yaml holds the superseded 2022 bands for 83 different "
        "indicators."
    )
    export.notes.append(
        "the export's own composite is a within-year min-max rescaling across the panel, so it is a "
        "relative ranking rather than an absolute score. It is read as a cross-check only."
    )
    return export


def sheet_for_economy(config, code: str) -> str | None:
    """The export sheet carrying one economy, or None where the panel does not include it.

    France is the live case: the export's 16-economy panel has no France sheet. Returning None rather than
    falling back to the euro area is deliberate, because a substituted economy in a comparative table is
    wrong in a way a reader cannot see.
    """
    from macrofield.model.honi import load_indicator_definitions

    document = load_indicator_definitions(Path(config.get("honi.indicators_config")))
    mapping = document.get("sheet_by_economy") or {}
    sheet = mapping.get(code)
    return str(sheet) if sheet else None


def score_from_export(
    config,
    code: str,
    export: Export | None = None,
    period: int | None = None,
) -> Any:
    """Score one economy from the export, for one year, using the programme's own aggregation.

    The publisher supplies the per-indicator scores; everything above that is ours: the dimension
    aggregation, the coverage floors, the weighting, and the four-stage taxonomy. See the module docstring
    for why the scoring step is not re-derived.

    Args:
        config: A loaded Config.
        code: This programme's economy code, for example `us`.
        export: A previously loaded export, so a comparative run reads the workbook once.
        period: The year to score. Defaults to the latest the export carries.

    Returns:
        A `macrofield.model.honi.HoNIScore`.

    Raises:
        HoNIExportError: If the economy is not in the export's panel.
    """
    from macrofield.model.honi import (
        IndicatorScore,
        aggregate_composite,
        classify_stage,
        load_indicator_definitions,
        score_dimension,
    )
    from macrofield.model.honi import HoNIScore

    sheet = sheet_for_economy(config, code)
    if sheet is None:
        raise HoNIExportError(
            f"the HoNI export's panel does not include {code!r}, so no reading can be produced for it. "
            f"See sheet_by_economy in config/honi_indicators.yaml."
        )

    export = export if export is not None else load_export(config)
    reading = export.for_sheet(sheet)
    year = int(period) if period is not None else reading.latest_year

    document = load_indicator_definitions(Path(config.get("honi.indicators_config")))
    definitions = document.get("indicators") or {}
    dimension_map = document.get("dimensions") or {}

    scores: dict[str, IndicatorScore] = {}
    for key, definition in definitions.items():
        value = reading.scores.get(key, {}).get(year)
        raw = reading.raw.get(key, {}).get(year)
        notes: list[str] = []
        if value is None:
            notes.append(f"no score in the export for {year}")
        scores[key] = IndicatorScore(
            key=key,
            name=str(definition.get("export_name", key)),
            score=value,
            raw_value=raw,
            notes=tuple(notes),
        )

    dimension_floor = float(config.get("honi.coverage_floor.dimension"))
    composite_floor = float(config.get("honi.coverage_floor.composite"))

    dimensions = {}
    for dimension_key, definition in dimension_map.items():
        dimensions[dimension_key] = score_dimension(
            key=dimension_key,
            label=str(definition.get("label", dimension_key)),
            indicator_keys=definition.get("indicators", []),
            scores=scores,
            coverage_floor=dimension_floor,
        )

    weights = config.get("honi.weights")
    composite, composite_notes = aggregate_composite(
        dimensions,
        {k: float(weights[k]) for k in dimensions if k in weights},
        composite_floor,
    )

    stages = config.get("honi.stages")
    stage = classify_stage(composite, stages) if composite is not None else None

    notes = list(export.notes) + list(reading.notes) + list(composite_notes)

    # The export's own composite, as a cross-check. It will not agree with ours and is not meant to.
    published = reading.published_composite.get(year)
    if published is not None and composite is not None:
        notes.append(
            f"the export's own composite for {year} is {published:.3f} against {composite:.3f} computed "
            f"here. They are not expected to agree: the export rescales each year's panel onto the 1 to 5 "
            f"axis, so its value is a within-year ranking, while this one is an absolute score on the "
            f"fixed axis. See docs/MODEL_SPEC.md section 4."
        )

    return HoNIScore(
        economy=code,
        period=year,
        dimensions=dimensions,
        composite=composite,
        stage=stage,
        capital_saturation=None,
        in_balanced_band=None,
        direction=str(config.get("honi.scale.direction")),
        strategy=str(config.get("honi.aggregation.strategy")),
        notes=notes,
    )


def _read_composite_sheet(workbook) -> dict[str, dict[int, float]]:
    """Read the published composite per economy per year."""
    if COMPOSITE_SHEET not in workbook.sheetnames:
        return {}
    sheet = workbook[COMPOSITE_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}

    header = rows[0]
    columns = {
        index: str(name).strip()
        for index, name in enumerate(header)
        if name is not None and str(name).strip()
    }
    out: dict[str, dict[int, float]] = {name: {} for name in columns.values()}
    for row in rows[1:]:
        year = _year(row[0] if row else None)
        if year is None:
            continue
        for index, name in columns.items():
            if index < len(row):
                value = _number(row[index])
                if value is not None:
                    out[name][year] = value
    return out


def _read_country_sheet(
    sheet,
    name: str,
    definitions: Mapping[str, Mapping[str, Any]],
    dimensions: Mapping[str, Mapping[str, Any]],
) -> EconomyReading:
    """Read one country sheet: the scored block, the raw block beneath it, and the sub-indices."""
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < FIRST_SCORE_ROW:
        raise HoNIExportError(f"sheet {name!r} is empty")

    header = rows[0]

    # Verify the recorded column map against the headers actually present. A mismatch would silently
    # attribute one indicator's scores to another, which is the worst failure this reader can have.
    mismatches: list[str] = []
    for key, definition in definitions.items():
        column = int(definition["export_column"]) - 1
        expected = str(definition["export_name"]).strip()
        found = (
            str(header[column]).strip()
            if column < len(header) and header[column] is not None
            else ""
        )
        if found != expected:
            mismatches.append(f"column {column + 1}: expected {expected!r}, found {found!r}")
    if mismatches:
        raise HoNIExportError(
            f"sheet {name!r} does not match the recorded column map in config/honi_indicators.yaml:\n  "
            + "\n  ".join(mismatches)
            + "\nCorrect the config rather than this reader: the map is configuration precisely so that "
            "an export whose columns move can be followed in one place."
        )

    # The scored block runs from FIRST_SCORE_ROW while column A holds a year.
    score_rows: list[tuple[int, tuple]] = []
    index = FIRST_SCORE_ROW - 1
    while index < len(rows):
        year = _year(rows[index][0] if rows[index] else None)
        if year is None:
            break
        score_rows.append((year, rows[index]))
        index += 1

    if not score_rows:
        raise HoNIExportError(f"sheet {name!r} carries no scored rows")

    years = [year for year, _ in score_rows]
    reading = EconomyReading(sheet=name, years=years)

    for key, definition in definitions.items():
        column = int(definition["export_column"]) - 1
        series: dict[int, float] = {}
        for year, row in score_rows:
            if column < len(row):
                value = _number(row[column])
                if value is not None:
                    series[year] = value
        reading.scores[key] = series

    for dimension_key, definition in dimensions.items():
        column = int(definition["export_column"]) - 1
        series = {}
        for year, row in score_rows:
            if column < len(row):
                value = _number(row[column])
                if value is not None:
                    series[year] = value
        reading.published_dimensions[dimension_key] = series

    # The raw block sits beneath the scored one, after a gap, in the same year order.
    raw_start = index + BLOCK_GAP
    raw_rows = []
    for offset in range(len(score_rows)):
        position = raw_start + offset
        if position >= len(rows):
            break
        raw_rows.append(rows[position])

    if len(raw_rows) == len(score_rows):
        for key, definition in definitions.items():
            column = int(definition["export_column"]) - 1
            series = {}
            for (year, _), row in zip(score_rows, raw_rows):
                if column < len(row):
                    value = _number(row[column])
                    if value is not None:
                        series[year] = value
            reading.raw[key] = series
    else:
        reading.notes.append(
            f"the raw-value block beneath the scores has {len(raw_rows)} rows against "
            f"{len(score_rows)} scored years, so raw values are not read for this economy. The scores "
            f"are unaffected."
        )

    return reading
