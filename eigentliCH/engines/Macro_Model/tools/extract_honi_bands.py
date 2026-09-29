"""Extract the HoNI scoring bands from the methodology workbook into config.

The bands are extracted rather than transcribed by hand so that the result is reproducible and its
provenance is recorded. Run this whenever the workbook is revised:

    .venv\\Scripts\\python.exe tools\\extract_honi_bands.py

Output: config/honi_bands.yaml, plus an integrity report on standard output.

Two decisions about how the workbook is read, both of which matter because the workbook contains
internal inconsistencies:

1. **The numeric cells are authoritative, not the prose interval labels.** Each score level carries
   both a text label (for example "[-1, -0.25[; ]3.75, 7]") and numeric boundary cells. Where the
   two disagree the cells are used, because the cells form a self-consistent contiguous chain across
   score levels while the labels do not. Every disagreement is reported.

2. **The scoring formulas in the workbook are not read at all.** The HoNI manual states that the IFS
   formula chains must be updated by hand whenever a band changes direction or moves position, which
   makes them a derived artefact that can fall out of step with the bands. This programme scores by
   interval membership against the bands directly, so the formulas are not a dependency.

Column layout of the Interpretation_Function sheet, one-based:

    1                       indicator or category name
    2, 7, 12, 17, 22        prose interval label for scores 1 to 5
    +1  lower               start of the lower interval
    +2  upper_lower_union   end of the lower interval, populated only for union bands
    +3  upper               start of the upper interval for union bands, else end of the interval
    +4  upper_upper_union   end of the upper interval, populated only for union bands
    27                      source
    28                      notes

So a simple band for score s covers [lower, upper]. A union band covers
[lower, upper_lower_union] together with [upper, upper_upper_union].
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

DEFAULT_WORKBOOK = Path(
    r"c:\Users\nicol\Desktop\SIM_NAS\Knowledge_Center\In progress"
    r"\Health of Nations 2023\Background\HoNI_Methodology.xlsx"
)
DEFAULT_OUTPUT = ROOT / "config" / "honi_bands.yaml"

SHEET = "Interpretation_Function"

#: One-based column at which each score level's block starts.
SCORE_START_COLUMN = {1: 2, 2: 7, 3: 12, 4: 17, 5: 22}
SOURCE_COLUMN = 27
NOTES_COLUMN = 28

#: A category header row, for example "Financial Economy - Level".
CATEGORY_PATTERN = re.compile(r"^(?P<name>.+?)\s*-\s*(?P<aspect>Level|Trend)\s*$", re.IGNORECASE)

#: The row at which the first table ends.
NOTES_MARKER = "scoring methodology notes"

#: Tolerance for comparing a label boundary against a cell boundary.
BOUNDARY_TOLERANCE = 1e-9


def parse_boundary(value: Any) -> float | None:
    """Return a numeric boundary, mapping the workbook's infinity glyphs to float infinities.

    Returns None for an empty cell, which is how a union column signals that it is unused.
    """
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    if text in {"∞", "+∞", "inf", "+inf"}:
        return math.inf
    if text in {"-∞", "-inf"}:
        return -math.inf
    # Some cells carry a stray comma or percent sign.
    cleaned = text.replace(",", "").replace("%", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def label_boundaries(label: Any) -> list[float]:
    """Return the numeric boundaries mentioned in a prose interval label, in order.

    Used only to cross-check the cells. Bracket direction is deliberately ignored, because the
    workbook is not consistent about it and the check is about the numbers.
    """
    if label is None:
        return []
    text = str(label)
    text = text.replace("∞", " INF ")
    found: list[float] = []
    for token in re.findall(r"-?\s*INF|-?\d+(?:\.\d+)?", text):
        stripped = token.replace(" ", "")
        if stripped.endswith("INF"):
            found.append(-math.inf if stripped.startswith("-") else math.inf)
        else:
            found.append(float(stripped))
    return found


@dataclass
class Indicator:
    """One HoNI indicator and its five score bands."""

    category: str
    aspect: str
    name: str
    row: int
    source: str | None
    notes: str | None
    #: score level -> list of [lower, upper] intervals
    bands: dict[int, list[list[float]]] = field(default_factory=dict)
    #: Integrity problems found while reading this indicator.
    problems: list[str] = field(default_factory=list)

    @property
    def key(self) -> str:
        """A stable unique key. Names repeat across categories, so the category is part of it."""
        slug = re.sub(r"[^a-z0-9]+", "_", self.name.lower()).strip("_")
        prefix = re.sub(r"[^a-z0-9]+", "_", f"{self.category}_{self.aspect}".lower()).strip("_")
        return f"{prefix}__{slug}"[:120]

    @property
    def direction(self) -> str:
        """Return the band direction, inferred from where score 1 sits relative to score 5.

        Reported for information. The scorer does not use it, because interval membership needs no
        direction, which is the point of scoring against the bands rather than the formulas.
        """
        if not self.bands.get(1) or not self.bands.get(5):
            return "unknown"
        if len(self.bands[1]) > 1 or len(self.bands[5]) > 1:
            return "union"
        return "positive" if self.bands[1][0][0] > self.bands[5][0][0] else "negative"


def read_indicators(workbook_path: Path) -> tuple[list[Indicator], list[str]]:
    """Read the indicator table from the workbook.

    Returns:
        The indicators, and a list of workbook-level integrity notes.
    """
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found in {workbook_path}")
    sheet = workbook[SHEET]
    rows = list(sheet.iter_rows(values_only=True))

    indicators: list[Indicator] = []
    notes: list[str] = []
    category = aspect = None
    end_row = None

    for index, row in enumerate(rows, start=1):
        first = row[0] if row else None
        text = str(first).strip() if first is not None else ""

        if text.lower().startswith(NOTES_MARKER):
            end_row = index
            break

        if not text or text.lower() == "indicator":
            continue

        match = CATEGORY_PATTERN.match(text)
        if match:
            category = match.group("name").strip()
            aspect = match.group("aspect").capitalize()
            continue

        if category is None:
            notes.append(f"row {index}: indicator {text!r} appears before any category header, skipped")
            continue

        indicator = Indicator(
            category=category,
            aspect=aspect or "Level",
            name=text,
            row=index,
            source=_clean(row[SOURCE_COLUMN - 1] if len(row) >= SOURCE_COLUMN else None),
            notes=_clean(row[NOTES_COLUMN - 1] if len(row) >= NOTES_COLUMN else None),
        )
        _read_bands(indicator, row)
        if indicator.bands:
            indicators.append(indicator)
        else:
            notes.append(f"row {index}: indicator {text!r} has no numeric bands, skipped")

    if end_row is not None:
        trailing = [
            index
            for index, row in enumerate(rows[end_row:], start=end_row + 1)
            if row and row[0] is not None and str(row[0]).strip()
        ]
        if trailing:
            notes.append(
                f"the sheet contains {len(trailing)} further non-empty rows after the notes block at "
                f"row {end_row} (first at row {trailing[0]}). That looks like a second, older copy of "
                f"the table. It is NOT extracted. Confirm which block is current before relying on "
                f"these bands."
            )

    return indicators, notes


def _clean(value: Any) -> str | None:
    """Return a stripped string, or None for an empty cell."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _read_bands(indicator: Indicator, row: tuple[Any, ...]) -> None:
    """Populate an indicator's bands from its row, recording any integrity problems."""
    for score, start in SCORE_START_COLUMN.items():
        def cell(offset: int) -> Any:
            position = start - 1 + offset
            return row[position] if len(row) > position else None

        label = cell(0)
        lower = parse_boundary(cell(1))
        upper_lower_union = parse_boundary(cell(2))
        upper = parse_boundary(cell(3))
        upper_upper_union = parse_boundary(cell(4))

        intervals: list[list[float]] = []
        if upper_upper_union is not None or upper_lower_union is not None:
            # A union band: two intervals.
            if lower is not None and upper_lower_union is not None:
                intervals.append([lower, upper_lower_union])
            if upper is not None and upper_upper_union is not None:
                intervals.append([upper, upper_upper_union])
            if not intervals:
                indicator.problems.append(
                    f"score {score}: union columns are populated but no complete interval could be "
                    f"read (lower={lower}, upper_lower_union={upper_lower_union}, upper={upper}, "
                    f"upper_upper_union={upper_upper_union})"
                )
        elif lower is not None and upper is not None:
            intervals.append([lower, upper])
        elif lower is not None or upper is not None:
            indicator.problems.append(
                f"score {score}: only one boundary is populated (lower={lower}, upper={upper}), so "
                f"the interval is incomplete"
            )

        for interval in intervals:
            if interval[0] > interval[1]:
                indicator.problems.append(
                    f"score {score}: interval [{interval[0]}, {interval[1]}] is inverted"
                )

        if intervals:
            indicator.bands[score] = intervals

        _cross_check_label(indicator, score, label, intervals)


def _cross_check_label(
    indicator: Indicator, score: int, label: Any, intervals: list[list[float]]
) -> None:
    """Compare the prose label's numbers against the numeric cells, reporting disagreements."""
    if label is None or not intervals:
        return
    from_label = label_boundaries(label)
    from_cells = [boundary for interval in intervals for boundary in interval]
    if len(from_label) != len(from_cells):
        indicator.problems.append(
            f"score {score}: label {str(label)!r} mentions {len(from_label)} boundaries but the "
            f"cells give {len(from_cells)}. Cells used."
        )
        return
    for position, (label_value, cell_value) in enumerate(zip(from_label, from_cells)):
        if math.isinf(label_value) and math.isinf(cell_value):
            if math.copysign(1, label_value) != math.copysign(1, cell_value):
                indicator.problems.append(
                    f"score {score}: label {str(label)!r} and cells disagree on the sign of an "
                    f"infinity at position {position}. Cells used."
                )
            continue
        if abs(label_value - cell_value) > BOUNDARY_TOLERANCE:
            indicator.problems.append(
                f"score {score}: label {str(label)!r} gives {label_value} at position {position} but "
                f"the cell gives {cell_value}. Cells used."
            )


def check_coverage(indicator: Indicator) -> list[str]:
    """Return coverage problems: overlapping score bands, or gaps in the real line.

    A value falling in a gap cannot be scored, and a value falling in an overlap could take two
    scores, so both are reported. The scorer handles them at run time as well, but catching them at
    extraction time is what makes them fixable in the workbook.
    """
    problems: list[str] = []
    flat = [
        (interval[0], interval[1], score)
        for score, intervals in sorted(indicator.bands.items())
        for interval in intervals
    ]
    flat.sort(key=lambda item: (item[0], item[1]))

    for (lo_a, hi_a, score_a), (lo_b, hi_b, score_b) in zip(flat, flat[1:]):
        if lo_b < hi_a - BOUNDARY_TOLERANCE:
            problems.append(
                f"scores {score_a} and {score_b} overlap: [{lo_a}, {hi_a}] and [{lo_b}, {hi_b}]"
            )
        elif lo_b > hi_a + BOUNDARY_TOLERANCE:
            problems.append(
                f"gap between scores {score_a} and {score_b}: [{lo_a}, {hi_a}] then [{lo_b}, {hi_b}]. "
                f"A value in ({hi_a}, {lo_b}) cannot be scored."
            )

    if flat:
        if not math.isinf(flat[0][0]):
            problems.append(
                f"bands are bounded below at {flat[0][0]}, so a value beneath it cannot be scored"
            )
        if not math.isinf(flat[-1][1]):
            problems.append(
                f"bands are bounded above at {flat[-1][1]}, so a value above it cannot be scored"
            )
    return problems


def build_document(
    indicators: list[Indicator], workbook_path: Path, workbook_notes: list[str]
) -> dict[str, Any]:
    """Assemble the YAML document, including provenance."""
    stat = workbook_path.stat()
    document: dict[str, Any] = {
        "provenance": {
            "source_document": "Health of Nations Index scoring model, methodology workbook",
            "source_path": str(workbook_path),
            "source_sheet": SHEET,
            "source_modified": dt.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "extracted_by": "tools/extract_honi_bands.py",
            "extracted_on": dt.date.today().isoformat(),
            # Machine-readable, because it is NOT the same as the current export's convention.
            # This 2022 workbook states "Key: 1 most healthy; 5 least healthy" in its own header, and
            # its bands bear that out (saving rates score 1 at [40, 100]; non-performing loans score 1
            # at [0, 2]). The 2026 export runs the other way: 5 is healthiest. The scorer converts
            # from this native direction into the canonical one, so the two cannot be mixed silently.
            "scale_direction": "lower_is_healthier",
            "scale": "1 most healthy, 5 least healthy, per the workbook's own header",
            "reading": (
                "numeric boundary cells are authoritative; prose interval labels are cross-checked "
                "and any disagreement is recorded in the indicator's problems list"
            ),
        },
        "workbook_notes": workbook_notes,
        "indicators": {},
    }

    for indicator in indicators:
        problems = list(indicator.problems) + check_coverage(indicator)
        document["indicators"][indicator.key] = {
            "name": indicator.name,
            "category": indicator.category,
            "aspect": indicator.aspect,
            "source_row": indicator.row,
            "source": indicator.source,
            "notes": indicator.notes,
            "direction": indicator.direction,
            "bands": {
                score: [[_yaml_number(lo), _yaml_number(hi)] for lo, hi in intervals]
                for score, intervals in sorted(indicator.bands.items())
            },
            "problems": problems,
        }
    return document


def _yaml_number(value: float) -> Any:
    """Render infinities as strings, because YAML has no portable infinity literal."""
    if math.isinf(value):
        return "-inf" if value < 0 else "inf"
    return value


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        print(f"workbook not found: {args.workbook}", file=sys.stderr)
        return 2

    indicators, workbook_notes = read_indicators(args.workbook)
    document = build_document(indicators, args.workbook, workbook_notes)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as handle:
        handle.write(
            "# Generated by tools/extract_honi_bands.py. Do not edit by hand: re-run the tool.\n"
            "# Scale: 1 most healthy, 5 least healthy.\n"
        )
        yaml.safe_dump(document, handle, sort_keys=False, allow_unicode=True, width=100)

    # Integrity report.
    by_category: dict[str, int] = {}
    flagged = 0
    for key, entry in document["indicators"].items():
        label = f"{entry['category']} - {entry['aspect']}"
        by_category[label] = by_category.get(label, 0) + 1
        if entry["problems"]:
            flagged += 1

    print(f"extracted {len(document['indicators'])} indicators to {args.output}")
    for label in sorted(by_category):
        print(f"  {by_category[label]:>3}  {label}")
    print(f"\n{flagged} indicators carry integrity problems:")
    for key, entry in document["indicators"].items():
        if entry["problems"]:
            print(f"\n  {entry['name']} (row {entry['source_row']}, {entry['category']})")
            for problem in entry["problems"]:
                print(f"    - {problem}")
    if workbook_notes:
        print("\nworkbook-level notes:")
        for note in workbook_notes:
            print(f"  - {note}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
