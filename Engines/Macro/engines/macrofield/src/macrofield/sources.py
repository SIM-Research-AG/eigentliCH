"""Raw source files to annual observations. Pure: bytes in, numbers out.

The parsers reproduce the ``macrofield`` connectors exactly, because the golden reconciliation
depends on it: the same file must yield the same float for the same year. They fetch nothing,
clean nothing and fill nothing. A missing value arrives as ``None`` and stays so.

Every series the engine can read is listed in :data:`CATALOGUE` under a stable id. The id is
what the store keys on and what the data-need specification names, so the Data Feed engine can
later supply the same ids and nothing downstream changes.
"""

from __future__ import annotations

import csv
import io
import json
import struct
import zipfile
from dataclasses import dataclass
from typing import Literal, Mapping, Optional

Annual = dict[int, Optional[float]]


class SourceError(ValueError):
    """A raw file does not have the shape its parser expects. The message says how."""


@dataclass(frozen=True)
class SeriesDef:
    series_id: str
    source: Literal["world_bank", "bis", "pwt", "bundesbank", "jst", "imf"]
    #: The identifier at source: an indicator code, a BIS key, a PWT column.
    identifier: str
    description: str
    units: str
    native_frequency: Literal["A", "Q", "M"]
    #: What the model uses it for, in the words of the data-need specification.
    index_block: str
    leading_concurrent_lagging: Literal["leading", "concurrent", "lagging"]
    #: True where a level error moves an economy against a threshold, not only a growth rate.
    scale_critical: bool


#: The BIS dimension codes selecting total credit to the non-financial sector (borrowing sector C,
#: not N, which is non-financial corporations only), from all lenders, at market value, as a
#: percentage of GDP, adjusted for breaks.
BIS_SELECTORS: Mapping[str, str] = {
    "TC_BORROWERS": "C",
    "TC_LENDERS": "A",
    "UNIT_TYPE": "770",
    "VALUATION": "M",
    "TC_ADJUST": "A",
}

CATALOGUE: tuple[SeriesDef, ...] = (
    SeriesDef("wb.NY.GDP.MKTP.CD", "world_bank", "NY.GDP.MKTP.CD",
              "Gross domestic product, current prices", "current US dollars", "A",
              "Y: output level every other quantity is expressed against", "concurrent", True),
    SeriesDef("wb.NE.GDI.FTOT.ZS", "world_bank", "NE.GDI.FTOT.ZS",
              "Gross fixed capital formation", "percent of GDP", "A",
              "K_R: perpetual-inventory extension of the capital ratio past PWT", "concurrent",
              False),
    SeriesDef("wb.NY.GDP.MKTP.KD.ZG", "world_bank", "NY.GDP.MKTP.KD.ZG",
              "Real GDP growth", "percent", "A",
              "K_R: perpetual-inventory extension of the capital ratio past PWT", "concurrent",
              False),
    SeriesDef("wb.NY.GNS.ICTR.ZS", "world_bank", "NY.GNS.ICTR.ZS",
              "Gross national savings", "percent of GDP", "A",
              "p_s: savings rate of the equations of motion", "concurrent", False),
    SeriesDef("wb.GC.NLD.TOTL.GD.ZS", "world_bank", "GC.NLD.TOTL.GD.ZS",
              "General government net lending (+) / borrowing (-)", "percent of GDP", "A",
              "S: stimulus proxy, sign reversed so a deficit is an injection", "leading", False),
    SeriesDef("wb.SP.POP.TOTL", "world_bank", "SP.POP.TOTL",
              "Population, total", "persons", "A",
              "p_b: population growth, which section 0.2 says p_b tracks", "lagging", False),
    SeriesDef("bis.total_credit", "bis", "WS_TC:C.A.M.770.A",
              "Total credit to the non-financial sector, all lenders, market value, adjusted "
              "for breaks", "percent of GDP", "Q",
              "Saturation axis and K_I (ratio times Y); also net new credit for S", "leading",
              True),
    SeriesDef("pwt.cn", "pwt", "cn",
              "Capital stock at current PPPs", "millions of 2017 US dollars", "A",
              "K_R: numerator of the capital-output ratio", "lagging", True),
    SeriesDef("pwt.cgdpo", "pwt", "cgdpo",
              "Output-side real GDP at current PPPs", "millions of 2017 US dollars", "A",
              "K_R: denominator of the capital-output ratio", "concurrent", True),
    SeriesDef("pwt.delta", "pwt", "delta",
              "Average depreciation rate of the capital stock", "rate", "A",
              "K_R: depreciation in the perpetual-inventory extension", "lagging", False),
    # Genreith's measure (Field Theory of Macroeconomics, 2014, section 2), for economies whose
    # saturation source is the bank balance sheet.
    SeriesDef("bbk.OU0308", "bundesbank", "BBBK1.M.OU0308",
              "Balance sheet total, all categories of banks (MFIs) in Germany; December value, "
              "DM converted at 1.95583 before 1999", "EUR bn", "M",
              "Saturation axis (Genreith's K): bank balance sheet over nominal GDP", "leading",
              True),
    SeriesDef("bbk.OU0115", "bundesbank", "BBBK1.M.OU0115",
              "Lending to domestic non-banks (non-MFIs), all categories of banks; December value, "
              "DM converted at 1.95583 before 1999", "EUR bn", "M",
              "Commercial-bank share (Genreith): loans over the bank balance sheet, Phase IV "
              "below 50 per cent", "leading", False),
    SeriesDef("wb.NY.GDP.MKTP.CN", "world_bank", "NY.GDP.MKTP.CN",
              "Gross domestic product, current prices, national currency", "LCU", "A",
              "Denominator of the bank balance-sheet ratio from the splice year on", "concurrent",
              True),
    # IMF DataMapper panels (calibration 1.3.0): where the World Bank or BIS publish nothing.
    SeriesDef("imf.GGXCNL_NGDP", "imf", "WEO:GGXCNL_NGDP",
              "General government net lending (+) / borrowing (-), IMF World Economic Outlook; "
              "forecast years cut at the calibration's last actual year", "percent of GDP", "A",
              "S: stimulus proxy where the World Bank series is absent or ends early", "leading",
              False),
    SeriesDef("imf.PVD_LS", "imf", "GDD:PVD_LS",
              "Private debt, loans and debt securities, IMF Global Debt Database",
              "percent of GDP", "A",
              "Saturation axis where BIS publishes none: private plus government debt, the same "
              "concept as BIS total credit to the non-financial sector", "leading", True),
    SeriesDef("imf.GG_DEBT_GDP", "imf", "GDD:GG_DEBT_GDP",
              "General government debt, IMF Global Debt Database", "percent of GDP", "A",
              "Government part of the IMF saturation axis", "lagging", True),
    SeriesDef("imf.GGXWDG_NGDP", "imf", "WEO:GGXWDG_NGDP",
              "General government gross debt, IMF World Economic Outlook", "percent of GDP", "A",
              "Government part of the IMF saturation axis where the Global Debt Database has "
              "none", "lagging", True),
    SeriesDef("jst.gdp", "jst", "gdp",
              "Nominal GDP, national currency (Jorda-Schularick-Taylor Macrohistory Database R6); "
              "for Germany West German DM before 1990", "LCU bn", "A",
              "Denominator of the bank balance-sheet ratio before the splice year (territory of "
              "the bank statistics)", "concurrent", True),
)

SERIES: Mapping[str, SeriesDef] = {s.series_id: s for s in CATALOGUE}


def _pairs_to_annual(pairs: list[tuple[int, Optional[float]]], label: str) -> Annual:
    """Keep the first observation of each year, sorted, as the old connectors did."""
    if not pairs:
        raise SourceError(f"no observations were parsed for {label}")
    out: Annual = {}
    for year, value in pairs:
        if year not in out:
            out[year] = value
    return dict(sorted(out.items()))


def _to_float(raw: object) -> Optional[float]:
    if raw is None:
        return None
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# World Bank indicators API (JSON)
# ---------------------------------------------------------------------------

def world_bank(payload: bytes, label: str) -> Annual:
    """One indicator for one country. ``None`` for a year the World Bank leaves empty."""
    try:
        document = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise SourceError(f"{label}: unparseable JSON") from exc
    if not isinstance(document, list) or len(document) < 2 or document[1] is None:
        message = ""
        if isinstance(document, list) and document and isinstance(document[0], dict):
            message = str(document[0].get("message", ""))
        raise SourceError(f"{label}: the World Bank returned no observations. {message}".strip())
    pairs = [(int(row["date"]), _to_float(row.get("value")))
             for row in document[1] if row.get("date") is not None]
    return _pairs_to_annual(pairs, label)


# ---------------------------------------------------------------------------
# BIS total credit, bulk CSV in a zip
# ---------------------------------------------------------------------------

def _code(cell: Optional[str]) -> str:
    """The code part of a BIS ``CODE: Label`` cell."""
    return "" if cell is None else str(cell).split(":", 1)[0].strip()


def bis_total_credit(payload: bytes, countries: Optional[set[str]] = None
                     ) -> tuple[dict[str, Annual], tuple[str, ...]]:
    """Total credit to the non-financial sector, % of GDP, fourth quarter, per borrower country.

    Credit is a stock, so the year-end quarter is taken rather than an average. Returns the
    series and the countries left out because the selectors matched more than one series for
    them (BIS added a dimension): an ambiguous country is excluded and named, never guessed.
    ``countries=None`` reads every borrower country in the file.
    """
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise SourceError("the BIS bulk download is not a zip archive") from exc
    names = [n for n in archive.namelist() if n.lower().endswith(".csv")]
    if not names:
        raise SourceError(f"the BIS archive holds no CSV: {archive.namelist()}")

    with archive.open(names[0]) as handle:
        reader = csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig"))
        if reader.fieldnames is None:
            raise SourceError("the BIS file has no header row")
        columns = {_code(f): f for f in reader.fieldnames}
        required = set(BIS_SELECTORS) | {"BORROWERS_CTY", "TIME_PERIOD", "OBS_VALUE"}
        missing = sorted(required - set(columns))
        if missing:
            raise SourceError(f"the BIS file lacks dimensions {missing}; its layout changed")
        variant = [h for c, h in sorted(columns.items())
                   if c in {"FREQ", "COLLECTION", "UNIT_MEASURE"}]

        pairs: dict[str, list[tuple[int, Optional[float]]]] = {}
        series_keys: dict[str, set[tuple[str, ...]]] = {}
        for row in reader:
            country = _code(row.get(columns["BORROWERS_CTY"]))
            if countries is not None and country not in countries:
                continue
            if any(_code(row.get(columns[d])) != v for d, v in BIS_SELECTORS.items()):
                continue
            period = str(row.get(columns["TIME_PERIOD"]) or "")
            if "-Q" not in period:
                continue
            year_text, quarter = period.split("-Q", 1)
            if not year_text.isdigit() or quarter.strip() != "4":
                continue
            raw = str(row.get(columns["OBS_VALUE"]) or "").strip()
            try:
                value = float(raw) if raw not in ("", "NaN") else None
            except ValueError:
                value = None
            pairs.setdefault(country, []).append((int(year_text), value))
            series_keys.setdefault(country, set()).add(tuple(_code(row.get(h)) for h in variant))

    ambiguous = tuple(sorted(c for c, k in series_keys.items() if len(k) > 1))
    # The old connector kept the last value written for a year; within one series a year's Q4
    # appears once, so first and last agree.
    data = {country: _pairs_to_annual(p, f"BIS total credit {country}")
            for country, p in pairs.items() if country not in ambiguous}
    return data, ambiguous


# ---------------------------------------------------------------------------
# Penn World Table 10.01 (xlsx). Read at load time only; needs openpyxl.
# ---------------------------------------------------------------------------

PWT_SHEET = "Data"
PWT_COLUMNS = ("cn", "cgdpo", "delta")


def penn_world_table(payload: bytes, countries: Optional[set[str]] = None
                     ) -> dict[str, dict[str, Annual]]:
    """``{country: {column: {year: value}}}``. ``countries=None`` reads every country."""
    from openpyxl import load_workbook  # load-time dependency, see pyproject [etl]

    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises several unrelated types
        raise SourceError(f"cannot open the Penn World Table workbook: {exc}") from exc
    if PWT_SHEET not in book.sheetnames:
        raise SourceError(f"the Penn World Table workbook has no {PWT_SHEET!r} sheet")
    rows = book[PWT_SHEET].iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    needed = ("countrycode", "year") + PWT_COLUMNS
    missing = [c for c in needed if c not in header]
    if missing:
        raise SourceError(f"the Penn World Table sheet lacks columns {missing}")
    at = {c: header.index(c) for c in needed}

    out: dict[str, dict[str, list[tuple[int, Optional[float]]]]] = {}
    for row in rows:
        country = row[at["countrycode"]]
        if country is None or row[at["year"]] is None:
            continue
        if countries is not None and country not in countries:
            continue
        year = int(row[at["year"]])
        per = out.setdefault(country, {c: [] for c in PWT_COLUMNS})
        for column in PWT_COLUMNS:
            per[column].append((year, _to_float(row[at[column]])))
    book.close()
    return {country: {column: _pairs_to_annual(p, f"PWT {column} {country}")
                      for column, p in cols.items()}
            for country, cols in out.items()}


# ---------------------------------------------------------------------------
# Deutsche Bundesbank time series (CSV from api.statistiken.bundesbank.de)
# ---------------------------------------------------------------------------

#: Irrevocable conversion rate of 1 January 1999. Bundesbank series are in DM before 1999 and in
#: euro from then on, so a series is converted to euro for every earlier period.
DM_PER_EUR = 1.95583


def bundesbank_year_end(payload: bytes, label: str) -> Annual:
    """A monthly Bundesbank series as its December value per year, in euro.

    Stocks are measured at year end, as BIS credit is (Q4). Years without a December
    observation are left out, never filled from another month.
    """
    text = payload.decode("utf-8-sig")
    pairs: list[tuple[int, Optional[float]]] = []
    for row in csv.reader(io.StringIO(text)):
        if len(row) < 2 or len(row[0]) != 7 or row[0][4] != "-":
            continue
        year_text, month_text = row[0][:4], row[0][5:]
        if not (year_text.isdigit() and month_text == "12"):
            continue
        value = _to_float(row[1]) if row[1].strip() not in ("", ".") else None
        year = int(year_text)
        if value is not None and year < 1999:
            value = value / DM_PER_EUR
        pairs.append((year, value))
    return _pairs_to_annual(pairs, label)


# ---------------------------------------------------------------------------
# Jorda-Schularick-Taylor Macrohistory Database (Stata 118 file)
# ---------------------------------------------------------------------------

def _tag(buf: bytes, name: str, start: int = 0) -> int:
    at = buf.find(f"<{name}>".encode(), start)
    if at < 0:
        raise SourceError(f"the Stata file lacks <{name}>")
    return at + len(name) + 2


def jst_column(payload: bytes, column: str, countries: Optional[set[str]] = None
               ) -> dict[str, Annual]:
    """One numeric column of the JST dataset per ISO3 country: ``{iso: {year: value}}``.

    A minimal reader for Stata's release 118 format, enough for the JST file (numeric columns,
    fixed-width strings); no third-party library. Stata's missing values arrive as ``None``.
    """
    if not payload.startswith(b"<stata_dta><header><release>118</release>"):
        raise SourceError("the JST file is not a Stata release 118 file")
    order = "<" if b"<byteorder>LSF</byteorder>" in payload[:200] else ">"
    k = struct.unpack(order + "H", payload[_tag(payload, "K"):_tag(payload, "K") + 2])[0]
    n = struct.unpack(order + "Q", payload[_tag(payload, "N"):_tag(payload, "N") + 8])[0]
    at = _tag(payload, "variable_types")
    types = struct.unpack(order + f"{k}H", payload[at:at + 2 * k])
    at = _tag(payload, "varnames")
    names = [payload[at + 129 * i: at + 129 * (i + 1)].split(b"\0", 1)[0].decode("utf-8")
             for i in range(k)]
    widths = {65526: 8, 65527: 4, 65528: 4, 65529: 2, 65530: 1, 32768: 8}
    size = [widths.get(t, t) for t in types]
    offsets = [sum(size[:i]) for i in range(k)]
    row_len = sum(size)
    for needed in ("iso", "year", column):
        if needed not in names:
            raise SourceError(f"the JST file has no column {needed!r}")
    iso_i, year_i, col_i = names.index("iso"), names.index("year"), names.index(column)
    if not (1 <= types[iso_i] <= 2045):
        raise SourceError("the JST iso column is not a fixed-width string")
    fmt = {65526: "d", 65527: "f", 65528: "l", 65529: "h", 65530: "b"}
    limits = {"d": 8.988465674311579e307, "f": 1.7014117e38, "l": 2147483620, "h": 32740,
              "b": 100}

    def number(row: bytes, i: int) -> Optional[float]:
        code = fmt.get(types[i])
        if code is None:
            raise SourceError(f"JST column {names[i]!r} is not numeric")
        value = struct.unpack(order + code, row[offsets[i]:offsets[i] + size[i]])[0]
        return None if value > limits[code] else float(value)

    data = _tag(payload, "data")
    out: dict[str, list[tuple[int, Optional[float]]]] = {}
    for r in range(n):
        row = payload[data + r * row_len: data + (r + 1) * row_len]
        iso = row[offsets[iso_i]:offsets[iso_i] + size[iso_i]].split(b"\0", 1)[0].decode()
        if countries is not None and iso not in countries:
            continue
        year = number(row, year_i)
        if year is None:
            continue
        out.setdefault(iso, []).append((int(year), number(row, col_i)))
    return {iso: _pairs_to_annual(p, f"JST {column} {iso}") for iso, p in out.items()}


# ---------------------------------------------------------------------------
# IMF DataMapper API (JSON panels)
# ---------------------------------------------------------------------------

def imf_datamapper(payload: bytes, label: str) -> dict[str, Annual]:
    """``{iso3: {year: value}}`` from a DataMapper indicator panel. Aggregates are kept too."""
    try:
        document = json.loads(payload)
        panels = {k: v for k, v in document["values"].items() if k and isinstance(v, dict)}
        (indicator, panel), = panels.items()
    except (json.JSONDecodeError, KeyError, ValueError) as exc:
        raise SourceError(f"{label}: not a single DataMapper indicator panel") from exc
    out: dict[str, Annual] = {}
    for area, values in panel.items():
        pairs = [(int(y), _to_float(v)) for y, v in values.items() if str(y).isdigit()]
        if pairs:
            out[area] = _pairs_to_annual(pairs, f"{label} {area}")
    return out
