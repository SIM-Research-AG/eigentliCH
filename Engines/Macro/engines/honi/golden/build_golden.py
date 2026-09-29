"""Freeze the MATLAB artefacts into self-contained JSON fixtures. Development only.

Run from the engine folder::

    python golden/build_golden.py               # rebuild the fixtures (datafeed running)
    python golden/build_golden.py --freeze-build  # also refreeze the build's own output

It reads the MATLAB *output* (``HoNI_Export.xlsx``, needs ``openpyxl`` from the ``dev``
extra) and asks datafeed for the *input*: the raw snapshot, over ``GET /panel`` like any
consumer. honi never reads source data itself. The tests read the JSON written here, so
they run anywhere.

Outputs, all under ``golden/``:

``matlab_export_2025-11.json``
    ``HoNI_Export.xlsx`` as MATLAB wrote it: raw indicator values, index scores, sector
    scores, national scores and capital saturation, years by countries.

``snapshot_2026-01-05/panel.json`` and ``source.json``
    datafeed's raw snapshot ``matlab-m_ts-2026-01-05`` (``panel@1.1.0``), exactly as MATLAB
    saw the data, and the datafeed checksum it must match.

``snapshot_2026-01-05/expected_1.0.0.json`` (``--freeze-build`` only)
    The current build's own output on that panel under calibration 1.0.0. A regression
    reference, not a reconciliation: see ``golden/README.md``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
MATLAB_DIR = Path(os.environ.get("HONI_MATLAB_DIR",
                                 r"C:\Users\nicol\Desktop\SIM_NAS\SIM_Tech\Master_Controller"))

SNAPSHOT_ID = "matlab-m_ts-2026-01-05"
SNAPSHOT_AS_OF = "2026-01-05"

#: MATLAB country field -> datafeed country code.
COUNTRIES = {
    "Brazil": "BR", "CH": "CH", "China": "CN", "EU": "EU", "India": "IN", "Indonesia": "ID",
    "Malaysia": "MY", "Philippines": "PH", "Thailand": "TH", "UK": "GB", "USA": "US",
    "Japan": "JP", "Bangladesh": "BD", "Vietnam": "VN", "Germany": "DE", "Spain": "ES",
}

#: HoNI_Export.xlsx column order (B..P) -> index name.
EXPORT_INDICES = (
    ("Budget_Balance", "budget_balance"), ("Monetary_Supply", "monetary_supply"),
    ("Gov_Debt", "government_debt"), ("Real_10y", "real_rate_10y"), ("Market_Cap", "market_cap"),
    ("Extdebt_afford", "external_debt_affordability"),
    ("Extdebt_expo", "external_debt_exposure"), ("Terms_Trade", "terms_of_trade"),
    ("Import_Res", "import_reserves"), ("Corruptopm", "corruption_freedom"),
    ("Cons_Power", "consumption_power"), ("Pop_Growth", "population_growth"),
    ("GDP_Growth_Capita", "gdp_per_capita_growth"), ("Depend_Cons", "consumption_dependency"),
    ("Labour_Force", "labour_force"),
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def cell(v) -> float | None:
    if v is None or (isinstance(v, float) and not np.isfinite(v)):
        return None
    return float(v)


def write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"wrote {path.relative_to(HERE.parent)} ({path.stat().st_size // 1024} KiB)")


# ---------------------------------------------------------------------------
# HoNI_Export.xlsx
# ---------------------------------------------------------------------------

def build_export() -> None:
    import openpyxl

    source = MATLAB_DIR / "HoNI_Export.xlsx"
    wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    head = list(wb["HoNI Score"].iter_rows(values_only=True))
    names = [n for n in head[0][1:] if n]
    years = [int(r[0]) for r in head[1:22]]
    national = [[cell(v) for v in r[1:1 + len(names)]] for r in head[1:22]]

    raw = {i: [] for _, i in EXPORT_INDICES}
    scores = {i: [] for _, i in EXPORT_INDICES}
    sectors = {"financial": [], "international": [], "real": []}
    saturation = []
    per_country = {}
    for name in names:
        rows = list(wb[name].iter_rows(values_only=True))
        header = rows[0]
        assert tuple(header[1:16]) == tuple(m for m, _ in EXPORT_INDICES), f"{name}: columns moved"
        assert (header[19], header[23], header[24], header[25]) == \
            ("Capital_Saturation", "Financial", "International", "Real"), f"{name}: layout moved"
        per_country[name] = {
            "scores": [[cell(v) for v in r[1:16]] for r in rows[1:22]],
            "raw": [[cell(v) for v in r[1:16]] for r in rows[24:45]],
            "sectors": [[cell(r[23]), cell(r[24]), cell(r[25])] for r in rows[1:22]],
            "saturation": [cell(r[19]) for r in rows[1:22]],
        }
    for t in range(len(years)):
        for k, (_, index) in enumerate(EXPORT_INDICES):
            raw[index].append([per_country[n]["raw"][t][k] for n in names])
            scores[index].append([per_country[n]["scores"][t][k] for n in names])
        for k, sector in enumerate(("financial", "international", "real")):
            sectors[sector].append([per_country[n]["sectors"][t][k] for n in names])
        saturation.append([per_country[n]["saturation"][t] for n in names])

    write(HERE / "matlab_export_2025-11.json", {
        "_source": str(source),
        "_source_sha256": sha256(source),
        "_note": ("HoNI_Export.xlsx as written by HoNI_Exporter.m / HoNI2.m v2.0 on 2025-11-13. "
                  "Year labels are MATLAB's hard-coded timeHoNI = 2004:2024. The M_TS.mat that "
                  "produced it no longer exists; see golden/README.md."),
        "years": years,
        "countries": [COUNTRIES[n] for n in names],
        "index_raw": raw,
        "index_scores": scores,
        "sectors": sectors,
        "national": national,
        "capital_saturation": saturation,
    })


# ---------------------------------------------------------------------------
# The input panel, from datafeed
# ---------------------------------------------------------------------------

def fetch_panel(datafeed_url: str) -> None:
    """Freeze datafeed's raw snapshot as honi's golden input, over HTTP like any consumer."""
    import httpx

    base = datafeed_url.rstrip("/")
    manifest = httpx.get(f"{base}/snapshots/{SNAPSHOT_ID}", timeout=60).raise_for_status().json()
    body = httpx.get(f"{base}/panel", params={"snapshot_id": SNAPSHOT_ID}, timeout=120).raise_for_status()
    folder = HERE / f"snapshot_{SNAPSHOT_AS_OF}"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "panel.json").write_bytes(body.content)
    write(folder / "source.json", {
        "_note": ("Frozen copy of datafeed's raw snapshot, fetched from GET /panel. The checksum "
                  "is datafeed's manifest checksum; test_golden checks the copy against it."),
        "snapshot_id": SNAPSHOT_ID, "datafeed_checksum": manifest["checksum"],
        "datafeed_source": manifest["source"],
    })
    print(f"wrote golden/snapshot_{SNAPSHOT_AS_OF}/panel.json from {base}")


def freeze_build() -> None:
    sys.path.insert(0, str(HERE.parent / "src"))
    from honi.calibration import DEFAULT
    from honi.contracts import Panel
    from honi.engine import as_matrix, run_model

    folder = HERE / f"snapshot_{SNAPSHOT_AS_OF}"
    panel = Panel.model_validate_json((folder / "panel.json").read_text(encoding="utf-8"))
    countries = list(COUNTRIES.values())
    out = run_model(panel, DEFAULT, countries, window=None, default_years=20)
    write(folder / f"expected_{DEFAULT.version}.json", {
        "_note": ("Frozen output of this build on panel.json under calibration "
                  f"{DEFAULT.version}. Regression only: it proves the build is stable, not "
                  "that it is right. Refreeze only with a reason, and put the reason in the "
                  "commit message and DECISIONS.md."),
        "calibration_version": DEFAULT.version,
        "countries": countries,
        "years": list(out.years),
        "national": as_matrix(out.national),
        "sectors": {k: as_matrix(v) for k, v in out.sectors.items()},
        "index_scores": {k: as_matrix(v) for k, v in out.index_scores.items()},
        "index_raw": {k: as_matrix(v) for k, v in out.index_raw.items()},
        "capital_saturation": as_matrix(out.capital_saturation),
    })


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--datafeed", default="http://127.0.0.1:8001",
                        help="datafeed to fetch the input panel from")
    parser.add_argument("--freeze-build", action="store_true",
                        help="also refreeze the build's own output (regression reference)")
    args = parser.parse_args(argv)
    if not MATLAB_DIR.is_dir():
        print(f"MATLAB sources not found at {MATLAB_DIR}; set HONI_MATLAB_DIR", file=sys.stderr)
        return 2
    build_export()
    fetch_panel(args.datafeed)
    if args.freeze_build:
        freeze_build()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
