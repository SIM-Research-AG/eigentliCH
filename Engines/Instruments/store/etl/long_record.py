"""Read the long annual record out of the Steiner (2021) workbooks.

This is offline ETL, not runtime. It is the only module in the project that imports
``openpyxl``, and nothing in ``api/`` or ``engines/`` may import it -- the engines read
the annual record from the database, so an engine run depends on a dated table rather than
on a spreadsheet that happens to be on somebody's disk.

**Which column of each workbook, and why it matters.** Most of these workbooks carry two
or three candidate vintages of the same series side by side -- a ``census`` column, a
``fred`` column -- and then a stitched ``*_cmbo`` column that joins them. The stitched
column is the series the reference implementation reads, and picking a different one
would silently change every figure downstream. The column indices below are transcribed
from ``1_data.mlx`` and each one is checked against the workbook's own header text on
load, so a re-ordered workbook fails here instead of producing a plausible wrong answer.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - offline tooling only
    raise ImportError(
        "store.etl.long_record needs openpyxl. It is an offline-ETL dependency; "
        "the engines and the API do not require it."
    ) from exc


@dataclass(frozen=True)
class SeriesSpec:
    """One column of one workbook, and what has to be true about it."""

    key: str
    filename: str
    column: int           # 1-based, as transcribed from the MATLAB source
    header_contains: str  # guard: the header cell must contain this, case-insensitively
    first_year: int
    last_year: int
    note: str


#: The register of source columns. ``column`` is 1-based and matches the MATLAB
#: ``table2array(T(:, n))`` index, because ``readtable`` consumes the header row and then
#: numbers the remaining columns from one, exactly as openpyxl does when we skip row 1.
SERIES: tuple[SeriesSpec, ...] = (
    # --- inputs to the eight economic indicators ---
    SeriesSpec("gdp", "gdp_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "Nominal GDP, stitched. Denominator of debt saturation."),
    SeriesSpec("gov_debt", "Gov_debt_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "Federal debt outstanding, stitched."),
    SeriesSpec("loans", "loans_1870-2020.xlsx", 6, "combo", 1870, 2020,
               "Bank loans outstanding, stitched. Note the header spells it 'combo'."),
    SeriesSpec("srate", "srate_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "1-year T-bill rate as a decimal. Short leg of the yield curve, and block 1."),
    SeriesSpec("lrate", "lrate_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "10-year Treasury rate as a decimal. Long leg of the yield curve."),
    SeriesSpec("narrow", "narrow_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "Narrow money stock. Monetary composition."),
    SeriesSpec("broad", "broad_1870-2020.xlsx", 5, "cmbo", 1870, 2020,
               "Broad money stock. Monetary composition."),
    SeriesSpec("bill", "bill_1870-2020.xlsx", 5, "cmbo", 1870, 2020,
               "Bill return. Financial-sentiment component and block 1's return."),
    SeriesSpec("bond", "bond_1870-2020.xlsx", 2, "mcrohist", 1870, 2020,
               "Long bond rate from the Macrohistory database. Block 2."),
    SeriesSpec("spx", "spx_1870-2020.xlsx", 2, "s&p", 1870, 2020,
               "S&P index level. Log-differenced for block 3."),
    SeriesSpec("erng", "erng_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "Wage earnings index, stitched."),
    SeriesSpec("ump", "ump_1870-2020.xlsx", 8, "cmbo", 1870, 2020,
               "Unemployment rate as a decimal, stitched."),
    SeriesSpec("cpi", "cpi_1870-2020.xlsx", 6, "cmbo", 1870, 2020,
               "CPI inflation rate as a decimal, stitched. Already a rate, not a level."),
    # --- the return blocks ---
    SeriesSpec("wheat", "wheat_1870-2020.xlsx", 4, "cmbo", 1870, 2020,
               "Wheat price. Half of the commodity index."),
    SeriesSpec("oil", "oil_1870-2020.xlsx", 4, "cmbo", 1870, 2020,
               "Oil price. The other half of the commodity index."),
    SeriesSpec("gold", "gold_1870-2020.xlsx", 2, "new york", 1870, 2020,
               "New York gold price per fine ounce. Log-differenced for block 6."),
    SeriesSpec("house", "house_1890-2020.xlsx", 5, "cmbo", 1890, 2020,
               "House price index, stitched. Truncated at 1890 -- block 8."),
    SeriesSpec("bondtr", "bondtr_1871-2015.xlsx", 2, "mchist", 1871, 2015,
               "Government bond total return, already a return. Truncated at 2015 -- block 4."),
    SeriesSpec("ag", "ag_1870-2020.xlsx", 4, "cmbo", 1870, 2020,
               "Farm business value. Sparse before 1910; the model slices from 1910 -- block 7."),
)

#: ``ag`` is the one series with holes. The reference implementation slices ``AG(41:end)``,
#: which is 1910 onwards, and that window is dense. Anything before it is dropped rather
#: than interpolated: a back-filled farm value would be an invention, and section 11.3 is
#: explicit that truncation is handled by matching only where the series exists.
AG_FIRST_DENSE_YEAR = 1910


class LongRecordError(ValueError):
    """Raised when a workbook does not look the way the register says it should."""


def _read_column(path: Path, spec: SeriesSpec) -> dict[int, float]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]  # MATLAB readtable also takes the first sheet
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise LongRecordError(f"{spec.filename} is empty")

    header = rows[0]
    if spec.column > len(header):
        raise LongRecordError(
            f"{spec.filename} has {len(header)} columns; the register wants column {spec.column}"
        )
    header_cell = str(header[spec.column - 1] or "")
    if spec.header_contains.lower() not in header_cell.lower():
        raise LongRecordError(
            f"{spec.filename} column {spec.column} is headed {header_cell!r}, which does not "
            f"contain {spec.header_contains!r}. The workbook layout has changed; the register "
            f"in store/etl/long_record.py must be re-checked against 1_data.mlx before loading."
        )

    out: dict[int, float] = {}
    for row in rows[1:]:
        if not row:
            continue
        year, value = row[0], row[spec.column - 1]
        if not isinstance(year, (int, float)):
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        out[int(year)] = float(value)
    return out


def _check_dense(key: str, values: dict[int, float], first: int, last: int) -> list[float]:
    missing = [y for y in range(first, last + 1) if y not in values]
    if missing:
        shown = ", ".join(str(y) for y in missing[:8])
        more = f" and {len(missing) - 8} more" if len(missing) > 8 else ""
        raise LongRecordError(
            f"series {key!r} has holes at {shown}{more}. The long record is never back-filled; "
            f"either the window is wrong or the workbook has changed."
        )
    return [values[y] for y in range(first, last + 1)]


@dataclass(frozen=True)
class RawSeries:
    """One source series as read, with the window it actually covers."""

    key: str
    first_year: int
    last_year: int
    values: tuple[float, ...]
    source_file: str
    source_column: int
    source_sha256: str
    note: str

    def year(self, y: int) -> float:
        if not self.first_year <= y <= self.last_year:
            raise KeyError(f"{self.key} does not cover {y}")
        return self.values[y - self.first_year]

    def window(self, first: int, last: int) -> list[float]:
        if first < self.first_year or last > self.last_year:
            raise KeyError(
                f"{self.key} covers {self.first_year}..{self.last_year}; "
                f"{first}..{last} was asked for"
            )
        return list(self.values[first - self.first_year : last - self.first_year + 1])


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_long_record(source_dir: Path) -> dict[str, RawSeries]:
    """Read every registered series out of ``source_dir``.

    Returns a mapping keyed by :attr:`SeriesSpec.key`. Each series is checked for density
    over its declared window and carries the SHA-256 of the workbook it came from, so a
    calibration can name the exact bytes it was produced against.
    """
    source_dir = Path(source_dir)
    if not source_dir.is_dir():
        raise LongRecordError(f"{source_dir} is not a directory")

    out: dict[str, RawSeries] = {}
    for spec in SERIES:
        path = source_dir / spec.filename
        if not path.is_file():
            raise LongRecordError(f"{spec.filename} not found in {source_dir}")
        values = _read_column(path, spec)
        first = AG_FIRST_DENSE_YEAR if spec.key == "ag" else spec.first_year
        series = _check_dense(spec.key, values, first, spec.last_year)
        out[spec.key] = RawSeries(
            key=spec.key,
            first_year=first,
            last_year=spec.last_year,
            values=tuple(series),
            source_file=spec.filename,
            source_column=spec.column,
            source_sha256=_sha256(path),
            note=spec.note,
        )
    return out


def iter_rows(record: dict[str, RawSeries]) -> Iterator[tuple[str, int, float]]:
    """Flatten the record into ``(series_key, year, value)`` rows for the store."""
    for key in sorted(record):
        series = record[key]
        for offset, value in enumerate(series.values):
            yield key, series.first_year + offset, value
